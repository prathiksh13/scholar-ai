import html
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.config import OUTPUT_DIR, SUPPORTED_FORMATS

FORMAT_STYLES = {
    "IEEE": {
        "font": "Times New Roman",
        "size": 10,
        "title_size": 24,
        "margin": 0.75,
        "line_spacing": 1.15,
        "citation": "[1]",
    },
    "ACM": {
        "font": "Linux Libertine",
        "size": 9,
        "title_size": 18,
        "margin": 1.0,
        "line_spacing": 1.0,
        "citation": "[Author et al. 2024]",
    },
    "Springer": {
        "font": "Times New Roman",
        "size": 10,
        "title_size": 20,
        "margin": 1.0,
        "line_spacing": 1.5,
        "citation": "(Author, 2024)",
    },
    "APA": {
        "font": "Times New Roman",
        "size": 12,
        "title_size": 14,
        "margin": 1.0,
        "line_spacing": 2.0,
        "citation": "(Author, 2024)",
    },
    "MLA": {
        "font": "Times New Roman",
        "size": 12,
        "title_size": 14,
        "margin": 1.0,
        "line_spacing": 2.0,
        "citation": "(Author 24)",
    },
}


def validate_format(fmt: str) -> str:
    fmt = fmt.upper()
    if fmt not in SUPPORTED_FORMATS:
        raise ValueError(f"Unsupported format. Choose from: {', '.join(SUPPORTED_FORMATS)}")
    return fmt


def text_to_html(text: str, fmt: str) -> str:
    fmt = validate_format(fmt)
    style = FORMAT_STYLES[fmt]
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if not lines:
        return "<p><em>Empty document</em></p>"

    blocks = []
    for i, line in enumerate(lines):
        escaped = html.escape(line)
        if i == 0:
            blocks.append(f'<h1 style="text-align:center">{escaped}</h1>')
        elif re.match(r"^(abstract|introduction|methodology|results|discussion|conclusion|references)\b", line, re.I):
            blocks.append(f"<h2>{escaped}</h2>")
        elif re.match(r"^\d+\.?\s+\w", line) or len(line) < 60:
            blocks.append(f"<h3>{escaped}</h3>")
        else:
            blocks.append(f"<p>{escaped}</p>")

    meta = (
        f'<p class="text-xs text-slate-500"><em>Formatted for {fmt} — '
        f'{style["font"]}, {style["size"]}pt, {style["line_spacing"]} spacing</em></p>'
    )
    return meta + "\n".join(blocks)


def apply_format_to_docx(source_text: str, document_id: str, fmt: str) -> Path:
    fmt = validate_format(fmt)
    style = FORMAT_STYLES[fmt]
    out_path = OUTPUT_DIR / f"{document_id}_{fmt}.docx"

    doc = Document()
    section = doc.sections[0]
    margin = Inches(style["margin"])
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin

    lines = [ln.strip() for ln in source_text.splitlines() if ln.strip()]
    if not lines:
        lines = ["Untitled Research Paper"]

    for i, line in enumerate(lines):
        if i == 0:
            p = doc.add_heading(line, level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(
            r"^(abstract|introduction|methodology|results|discussion|conclusion|references)\b",
            line,
            re.I,
        ):
            doc.add_heading(line, level=1)
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.line_spacing = style["line_spacing"]
            for run in p.runs:
                run.font.name = style["font"]
                run.font.size = Pt(style["size"])

    doc.save(str(out_path))
    return out_path


def pandoc_convert(input_path: Path, output_path: Path, fmt: str) -> bool:
    """Optional Pandoc enhancement for exports."""
    try:
        subprocess.run(
            [
                "pandoc",
                str(input_path),
                "-o",
                str(output_path),
                "--metadata",
                f"title=ScholarAI {fmt} Document",
            ],
            check=True,
            capture_output=True,
            timeout=60,
        )
        return output_path.exists()
    except (FileNotFoundError, subprocess.CalledProcessError):
        return False
