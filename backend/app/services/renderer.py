"""Semantic JSON -> DOCX/PDF renderer with real pagination and figures."""

from __future__ import annotations

import base64
import re
from dataclasses import dataclass
from html import unescape
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.config import OUTPUT_DIR
from app.models.document_schema import DocumentElement, Figure, SemanticDocument, TableBlock
from app.services.layout_engine import layout_document, render_layout_to_pdf
from app.rules.ieee import get_ieee_rules

try:  # optional dependency, installed in the backend environment
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        Image as RLImage,
        KeepTogether,
        NextPageTemplate,
        PageBreak,
        PageTemplate,
        Paragraph as RLParagraph,
        Spacer,
        Table as RLTable,
        TableStyle,
    )

    REPORTLAB_AVAILABLE = True
except Exception:  # pragma: no cover - fallback path
    REPORTLAB_AVAILABLE = False


@dataclass
class FlowBlock:
    kind: str
    order: int
    value: object


def render_semantic_to_docx(doc: SemanticDocument, document_id: str) -> Path:
    rules = get_ieee_rules()
    out = OUTPUT_DIR / f"{document_id}_ieee.docx"
    d = Document()

    margin = Inches(rules.margins_inch)
    for section in d.sections:
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin

    blocks = _ordered_body_blocks(doc)
    if not doc.elements:
        title_p = d.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(doc.title)
        title_run.bold = True
        title_run.font.name = rules.body_font
        title_run.font.size = Pt(16)

        if doc.authors:
            ap = d.add_paragraph("; ".join(a.name for a in doc.authors))
            ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_body(ap, rules.body_font, rules.body_size_pt)

        if doc.abstract:
            ah = d.add_paragraph()
            ah.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ar = ah.add_run("Abstract")
            ar.bold = True
            _style_body(ah, rules.body_font, rules.body_size_pt)
            ap = d.add_paragraph(doc.abstract)
            _style_body(ap, rules.body_font, rules.body_size_pt)

        if doc.keywords:
            kp = d.add_paragraph("Keywords: " + ", ".join(doc.keywords))
            _style_body(kp, rules.body_font, rules.body_size_pt)

    body_section = d.add_section(WD_SECTION.NEW_PAGE)
    body_section.top_margin = margin
    body_section.bottom_margin = margin
    body_section.left_margin = margin
    body_section.right_margin = margin
    _set_section_columns(body_section, 2, 0.22)

    _append_docx_blocks(d, blocks, rules.body_font, rules.body_size_pt)

    d.save(str(out))
    return out


def render_semantic_to_pdf(doc: SemanticDocument, document_id: str) -> Path:
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError("PDF export requires reportlab. Install the backend dependencies again.")

    out = OUTPUT_DIR / f"{document_id}_ieee.pdf"
    layout = layout_document(doc)
    render_layout_to_pdf(layout, str(out), title=doc.title)
    return out


def render_editor_html_to_docx(editor_html: str, document_id: str) -> Path:
    """Render the editable HTML back to a DOCX file."""
    rules = get_ieee_rules()
    out = OUTPUT_DIR / f"{document_id}_edited.docx"
    d = Document()

    for section in d.sections:
        margin = Inches(rules.margins_inch)
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin

    blocks = re.findall(
        r"(<h[1-6][^>]*>.*?</h[1-6]>|<p[^>]*>.*?</p>|<figure[^>]*>.*?</figure>|<ol[^>]*>.*?</ol>|<table[^>]*>.*?</table>)",
        editor_html,
        flags=re.I | re.S,
    )
    if not blocks:
        blocks = [editor_html]

    for block in blocks:
        _append_html_block(d, block, rules.body_font, rules.body_size_pt)

    d.save(str(out))
    return out


def _ordered_body_blocks(doc: SemanticDocument) -> list[FlowBlock]:
    if doc.elements:
        blocks: list[FlowBlock] = []
        for element in doc.elements:
            blocks.append(FlowBlock(element.type, element.order, element))
        return blocks

    blocks: list[FlowBlock] = []
    blocks.extend(FlowBlock("section", sec.source_index, sec) for sec in doc.sections)
    blocks.extend(FlowBlock("figure", fig.source_index, fig) for fig in doc.figures)
    blocks.extend(FlowBlock("table", tbl.source_index, tbl) for tbl in doc.tables)
    if doc.references:
        blocks.append(FlowBlock("references", len(blocks), None))
        blocks.extend(FlowBlock("reference", ref.order, ref) for ref in sorted(doc.references, key=lambda item: item.order))
    return sorted(blocks, key=lambda item: item.order)


def _front_matter_flowables(doc: SemanticDocument, styles: dict[str, ParagraphStyle], max_width: float):
    flowables = []
    flowables.append(RLParagraph(_safe_text(doc.title), styles["title"]))
    if doc.authors:
        authors = ", ".join(_safe_text(a.name) for a in doc.authors)
        flowables.append(Spacer(1, 0.08 * inch))
        flowables.append(RLParagraph(authors, styles["authors"]))
    if doc.abstract:
        flowables.append(Spacer(1, 0.12 * inch))
        flowables.append(RLParagraph("Abstract", styles["heading2"]))
        flowables.append(RLParagraph(_safe_text(doc.abstract), styles["body"]))
    if doc.keywords:
        flowables.append(Spacer(1, 0.08 * inch))
        flowables.append(RLParagraph("<b>Keywords:</b> " + ", ".join(_safe_text(k) for k in doc.keywords), styles["body"]))
    return flowables


def _body_flowables(doc: SemanticDocument, styles: dict[str, ParagraphStyle], max_width: float):
    return _blocks_to_flowables(_ordered_body_blocks(doc), styles, max_width, skip_front_matter=True)


def _blocks_to_flowables(blocks: list[FlowBlock], styles: dict[str, ParagraphStyle], max_width: float, skip_front_matter: bool):
    flowables = []
    reference_counter = 0
    for block in blocks:
        value = block.value
        if block.kind == "title":
            if isinstance(value, DocumentElement):
                flowables.append(RLParagraph(_safe_text(value.content), styles["title"]))
            continue
        if block.kind == "author":
            if isinstance(value, DocumentElement):
                flowables.append(RLParagraph(_safe_text(value.content), styles["authors"]))
            continue
        if block.kind == "keyword":
            if isinstance(value, DocumentElement):
                flowables.append(RLParagraph("<b>Keywords:</b> " + _safe_text(value.content), styles["body"]))
            continue
        if skip_front_matter and block.kind in {"title", "author", "keyword"}:
            continue
        if block.kind == "heading" and isinstance(value, DocumentElement) and value.content.lower().strip() == "abstract":
            flowables.append(Spacer(1, 0.04 * inch))
            flowables.append(RLParagraph(_safe_text(value.content), styles["heading2"]))
            continue
        if block.kind == "paragraph" and isinstance(value, DocumentElement) and value.content and not value.src:
            flowables.append(RLParagraph(_safe_text(value.content), styles["body"]))
            continue
        if block.kind == "page_break":
            flowables.append(PageBreak())
            continue
        if block.kind == "references":
            flowables.append(Spacer(1, 0.12 * inch))
            flowables.append(RLParagraph("References", styles["heading2"]))
            continue
        if block.kind == "reference":
            reference_counter += 1
            if hasattr(value, "text"):
                ref_label = getattr(value, "order", reference_counter) or reference_counter
                flowables.append(RLParagraph(f"[{ref_label}] {_safe_text(value.text)}", styles["body"]))
            elif isinstance(value, DocumentElement):
                ref_label = reference_counter if reference_counter else value.order + 1
                flowables.append(RLParagraph(f"[{ref_label}] {_safe_text(value.content)}", styles["body"]))
            continue
        if block.kind == "heading" or block.kind == "paragraph":
            if isinstance(value, DocumentElement):
                if value.content:
                    style_name = f"heading{min(value.level + 1, 3)}" if block.kind == "heading" else "body"
                    if block.kind == "heading":
                        flowables.append(Spacer(1, 0.04 * inch))
                    flowables.append(RLParagraph(_safe_text(value.content), styles[style_name]))
            else:
                sec = value
                if sec.title:
                    flowables.append(Spacer(1, 0.05 * inch))
                    flowables.append(RLParagraph(_safe_text(sec.title), styles[f"heading{min(sec.level + 1, 3)}"]))
                if sec.content:
                    flowables.append(RLParagraph(_safe_text(sec.content), styles["body"]))
            continue
        if block.kind in {"image", "figure"}:
            if isinstance(value, DocumentElement):
                fig = _figure_from_document_element(value)
                flowables.extend(_figure_flowables(fig, styles, max_width))
            else:
                flowables.extend(_figure_flowables(value, styles, max_width))
            continue
        if block.kind == "section":
            sec = value
            if sec.title:
                flowables.append(Spacer(1, 0.05 * inch))
                flowables.append(RLParagraph(_safe_text(sec.title), styles[f"heading{min(sec.level + 1, 3)}"]))
            if sec.content:
                flowables.append(RLParagraph(_safe_text(sec.content), styles["body"]))
        elif block.kind == "figure":
            flowables.extend(_figure_flowables(value, styles, max_width))
        elif block.kind == "table":
            if isinstance(value, DocumentElement):
                tbl = _table_from_document_element(value)
                flowables.extend(_table_flowables(tbl, styles, max_width))
            else:
                flowables.extend(_table_flowables(value, styles, max_width))
    return flowables


def _split_front_matter_blocks(blocks: list[FlowBlock]) -> tuple[list[FlowBlock], list[FlowBlock]]:
    front: list[FlowBlock] = []
    body: list[FlowBlock] = []
    in_front = True
    for block in blocks:
        value = block.value
        if in_front:
            front.append(block)
            if block.kind == "heading" and isinstance(value, DocumentElement):
                heading_text = value.content.strip().lower()
                if heading_text not in {"abstract", "keywords"}:
                    in_front = False
            elif block.kind in {"section", "table", "figure", "image", "references"}:
                in_front = False
        else:
            body.append(block)
    if not body and front:
        return front, body
    return front, body


def _figure_flowables(fig: Figure, styles: dict[str, ParagraphStyle], max_width: float):
    items = []
    if fig.image_b64:
        image_bytes = base64.b64decode(fig.image_b64)
        reader = ImageReader(BytesIO(image_bytes))
        img_w, img_h = reader.getSize()
        draw_w = min(max_width, img_w)
        draw_h = draw_w * (img_h / img_w)
        image = RLImage(BytesIO(image_bytes), width=draw_w, height=draw_h)
        image.hAlign = "CENTER"
        items.append(image)
    else:
        items.append(RLParagraph(_safe_text(f"Figure {fig.id}"), styles["body"]))
    items.append(Spacer(1, 0.05 * inch))
    items.append(RLParagraph(_safe_text(fig.caption or fig.id), styles["caption"]))
    return [KeepTogether(items)]


def _figure_from_document_element(element: DocumentElement) -> Figure:
    src = element.src or ""
    image_b64 = None
    content_type = None
    if src.startswith("data:") and "," in src:
        head, image_b64 = src.split(",", 1)
        content_type = head[5:].split(";", 1)[0]
    return Figure(
        id=f"img-{element.order}",
        caption=element.caption or element.content or f"Figure {element.order}",
        width_pct=100.0,
        position="inline",
        alt=element.caption or element.content or f"Figure {element.order}",
        image_b64=image_b64,
        content_type=content_type,
        source_index=element.order,
    )


def _table_flowables(tbl: TableBlock, styles: dict[str, ParagraphStyle], max_width: float):
    if not tbl.rows:
        return []
    col_count = max(len(row) for row in tbl.rows)
    col_width = max_width / max(col_count, 1)
    data = []
    for row in tbl.rows:
        cells = [_safe_text(cell) for cell in row]
        cells += [""] * (col_count - len(cells))
        data.append(cells)
    table = RLTable(data, colWidths=[col_width] * col_count, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "Times-Roman"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("LEADING", (0, 0), (-1, -1), 11),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5eef8")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#0f172a")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#93c5fd")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    caption = RLParagraph(_safe_text(tbl.caption or tbl.id), styles["caption"])
    return [KeepTogether([table, Spacer(1, 0.04 * inch), caption])]


def _table_from_document_element(element: DocumentElement) -> TableBlock:
    rows = _html_table_rows(element.html or "")
    return TableBlock(
        id=f"tbl-{element.order}",
        caption=element.caption or element.content or f"Table {element.order}",
        rows=rows,
        width_pct=100.0,
        source_index=element.order,
        html=element.html,
    )


def _append_docx_block(doc: Document, block: FlowBlock, font_name: str, size_pt: int) -> None:
    if block.kind == "page_break":
        doc.add_page_break()
        return

    if block.kind == "title":
        element = block.value
        if isinstance(element, DocumentElement):
            p = doc.add_paragraph(element.content)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.runs[0] if p.runs else p.add_run(element.content)
            run.bold = True
            _style_body(p, font_name, max(size_pt + 6, 14))
        return

    if block.kind == "author":
        element = block.value
        if isinstance(element, DocumentElement):
            p = doc.add_paragraph(element.content)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_body(p, font_name, size_pt)
        return

    if block.kind == "keyword":
        element = block.value
        if isinstance(element, DocumentElement):
            p = doc.add_paragraph(f"Keywords: {element.content}")
            _style_body(p, font_name, size_pt)
        return

    if block.kind == "heading" or block.kind == "paragraph" or block.kind == "caption":
        element = block.value
        if isinstance(element, DocumentElement):
            if element.content:
                if block.kind == "heading":
                    doc.add_heading(element.content, level=min(element.level + 1, 3))
                elif block.kind == "caption":
                    p = doc.add_paragraph(element.content)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _style_body(p, font_name, max(8, size_pt - 1))
                else:
                    p = doc.add_paragraph(element.content)
                    _style_body(p, font_name, size_pt)
                return

    if block.kind == "section":
        sec = block.value
        if sec.title:
            doc.add_heading(sec.title, level=min(sec.level + 1, 3))
        if sec.content:
            p = doc.add_paragraph(sec.content)
            _style_body(p, font_name, size_pt)
        return

    if block.kind in {"figure", "image"}:
        if isinstance(block.value, DocumentElement):
            fig = _figure_from_document_element(block.value)
        else:
            fig = block.value
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if fig.image_b64:
            image_bytes = base64.b64decode(fig.image_b64)
            run = p.add_run()
            run.add_picture(BytesIO(image_bytes), width=Inches(3.1))
        else:
            p.add_run(f"[Figure {fig.id}]")
        cap = doc.add_paragraph(fig.caption or fig.id)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_body(cap, font_name, max(8, size_pt - 1))
        return

    if block.kind == "table":
        tbl_value = block.value
        if isinstance(tbl_value, DocumentElement) and tbl_value.html:
            _append_table_html_to_docx(doc, tbl_value.html, font_name, size_pt)
            return
        tbl: TableBlock = tbl_value
        if tbl.rows:
            table = doc.add_table(rows=len(tbl.rows), cols=max(1, max(len(r) for r in tbl.rows)))
            table.style = "Table Grid"
            for ri, row in enumerate(tbl.rows):
                for ci, cell in enumerate(row):
                    table.rows[ri].cells[ci].text = cell
        cap = doc.add_paragraph(tbl.caption or tbl.id)
        _style_body(cap, font_name, max(8, size_pt - 1))
        return

    if block.kind == "references":
        doc.add_heading("References", level=1)
        return

    if block.kind == "image":
        element = block.value
        if isinstance(element, DocumentElement):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if element.src and element.src.startswith("data:"):
                _append_docx_image_from_src(doc, p, element.src)
            elif element.html:
                image_src = _extract_img_src(element.html)
                if image_src:
                    _append_docx_image_from_src(doc, p, image_src)
            if element.caption:
                cap = doc.add_paragraph(element.caption)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _style_body(cap, font_name, max(8, size_pt - 1))
            return


def _append_docx_blocks(doc: Document, blocks: list[FlowBlock], font_name: str, size_pt: int) -> None:
    for block in blocks:
        _append_docx_block(doc, block, font_name, size_pt)


def _append_html_block(doc: Document, block: str, font_name: str, size_pt: int) -> None:
    if re.match(r"<h1", block, re.I):
        doc.add_heading(_strip_html(block), level=0)
        return
    if re.match(r"<h2", block, re.I):
        doc.add_heading(_strip_html(block), level=1)
        return
    if re.match(r"<h3", block, re.I):
        doc.add_heading(_strip_html(block), level=2)
        return
    if block.lstrip().lower().startswith("<figure"):
        caption = _extract_tag_text(block, "figcaption") or _strip_html(block)
        image_src = _extract_img_src(block)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if image_src:
            _append_docx_image_from_src(doc, p, image_src)
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_body(cap, font_name, max(8, size_pt - 1))
        return
    if block.lstrip().lower().startswith("<ol"):
        for item in re.findall(r"<li[^>]*>(.*?)</li>", block, flags=re.I | re.S):
            p = doc.add_paragraph(_strip_html(item), style="List Number")
            _style_body(p, font_name, size_pt)
        return
    if block.lstrip().lower().startswith("<table"):
        rows = re.findall(r"<tr[^>]*>(.*?)</tr>", block, flags=re.I | re.S)
        if rows:
            cols = max(1, max(len(re.findall(r"<t[dh][^>]*>.*?</t[dh]>", row, flags=re.I | re.S)) for row in rows))
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for row_index, row_html in enumerate(rows):
                cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row_html, flags=re.I | re.S)
                for cell_index, cell_html in enumerate(cells):
                    table.rows[row_index].cells[cell_index].text = _strip_html(cell_html)
        return
    text = _strip_html(block)
    if text:
        p = doc.add_paragraph(text)
        _style_body(p, font_name, size_pt)


def _append_docx_image_from_src(doc: Document, paragraph, src: str) -> None:
    if src.startswith("data:image/") and ";base64," in src:
        encoded = src.split(",", 1)[1]
        try:
            paragraph.add_run().add_picture(BytesIO(base64.b64decode(encoded)), width=Inches(3.1))
        except Exception:
            paragraph.add_run("[image]")
    else:
        paragraph.add_run("[image]")


def _append_table_html_to_docx(doc: Document, html: str, font_name: str, size_pt: int) -> None:
    rows = _html_table_rows(html)
    if rows:
        table = doc.add_table(rows=len(rows), cols=max(1, max(len(row) for row in rows)))
        table.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci, cell in enumerate(row):
                table.rows[ri].cells[ci].text = cell


def _html_table_rows(html: str) -> list[list[str]]:
    rows: list[list[str]] = []
    try:
        root = etree.HTML(html)
        if root is None:
            return rows
        for tr in root.xpath('.//tr'):
            row: list[str] = []
            for cell in tr.xpath('./th|./td'):
                text = ''.join(cell.xpath('.//text()')).strip()
                row.append(_strip_html(text))
            if row:
                rows.append(row)
    except Exception:
        return rows
    return rows


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "IEEETitle",
            parent=base["Title"],
            fontName="Times-Bold",
            fontSize=16,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=6,
        ),
        "authors": ParagraphStyle(
            "IEEEAuthors",
            parent=base["BodyText"],
            fontName="Times-Roman",
            fontSize=10,
            leading=12,
            alignment=TA_CENTER,
            spaceAfter=6,
        ),
        "heading2": ParagraphStyle(
            "IEEEH2",
            parent=base["Heading2"],
            fontName="Times-Bold",
            fontSize=11,
            leading=13,
            alignment=TA_LEFT,
            spaceBefore=6,
            spaceAfter=4,
        ),
        "heading3": ParagraphStyle(
            "IEEEH3",
            parent=base["Heading3"],
            fontName="Times-Bold",
            fontSize=10.5,
            leading=12,
            alignment=TA_LEFT,
            spaceBefore=4,
            spaceAfter=2,
        ),
        "body": ParagraphStyle(
            "IEEBody",
            parent=base["BodyText"],
            fontName="Times-Roman",
            fontSize=10,
            leading=12,
            alignment=TA_JUSTIFY,
            spaceAfter=4,
        ),
        "caption": ParagraphStyle(
            "IEEECaption",
            parent=base["BodyText"],
            fontName="Times-Italic",
            fontSize=8,
            leading=9,
            alignment=TA_CENTER,
            spaceAfter=4,
        ),
    }


def _draw_pdf_page(canvas, doc):
    canvas.saveState()
    canvas.setFont("Times-Roman", 8)
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 0.45 * inch, str(canvas.getPageNumber()))
    canvas.restoreState()


def _set_section_columns(section, count: int, gap_inch: float) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(int(gap_inch * 1440)))


def _style_body(paragraph, font_name: str, size_pt: int) -> None:
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(size_pt)


def _safe_text(value: str) -> str:
    return unescape(re.sub(r"<[^>]+>", "", value)).strip()


def _strip_html(value: str) -> str:
    return _safe_text(value)


def _extract_tag_text(value: str, tag: str) -> str:
    match = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", value, flags=re.I | re.S)
    return _strip_html(match.group(1)) if match else ""


def _extract_img_src(value: str) -> str | None:
    match = re.search(r'<img[^>]+src=["\']([^"\']+)["\']', value, flags=re.I | re.S)
    return match.group(1) if match else None